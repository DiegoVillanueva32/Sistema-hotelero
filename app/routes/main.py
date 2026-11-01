import io
import os
import hmac

from flask import Blueprint, Response, flash, redirect, render_template, request, session, url_for

from app.persistence.persistencia import validar_password_admin
from app.services.hotel_service import hotel_service
from app.timezone import hora_local

bp = Blueprint("main", __name__)

RECEPCION_PASSWORD = os.getenv("HOTEL_RECEPCION_PASSWORD", "recep_hotel99")


@bp.before_app_request
def exigir_inicio_sesion():
    endpoint = request.endpoint or ""
    if endpoint in {"main.login", "static"}:
        return None
    if session.get("recepcion_autorizado"):
        return None
    return redirect(url_for("main.login", next=request.path))


@bp.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        recepcionista = request.form.get("recepcionista", "").strip()
        password = request.form.get("password", "")
        if not recepcionista:
            flash("Ingresa el nombre del recepcionista.", "error")
        elif not hmac.compare_digest(password, RECEPCION_PASSWORD):
            flash("Contraseña incorrecta.", "error")
        else:
            session.clear()
            session["recepcion_autorizado"] = True
            session["recepcionista"] = recepcionista
            session.permanent = False
            destino = request.args.get("next")
            if (
                not destino
                or not destino.startswith("/")
                or destino.startswith("//")
                or destino == "/login"
                or destino == "/run.py"
            ):
                destino = url_for("main.dashboard")
            return redirect(destino)

    return render_template("login.html")


@bp.route("/logout")
def logout():
    session.clear()
    flash("Sesión cerrada correctamente.", "success")
    return redirect(url_for("main.login"))


def ejecutar_accion(accion, *args, **kwargs):
    try:
        resultado = accion(*args, **kwargs)
        if isinstance(resultado, tuple):
            ok, mensaje = resultado
            flash(mensaje, "success" if ok else "error")
            if ok:
                hotel_service.persistir()
        else:
            if resultado:
                flash(str(resultado), "success")
            hotel_service.persistir()
        return True
    except Exception as exc:
        flash(f"No se pudo completar la operación: {exc}", "error")
        return False


@bp.route("/")
def dashboard():
    gestor = hotel_service.gestor
    habitaciones = list(gestor.consultar_matriz().values())
    estados = {}
    for habitacion in habitaciones:
        estados[habitacion.estado] = estados.get(habitacion.estado, 0) + 1
    return render_template(
        "dashboard.html",
        total=len(habitaciones),
        estados=estados,
        ocupadas=sum(1 for h in habitaciones if h.estado == "Ocupado"),
        reservas=len(gestor.historial_reservas),
        pagos=sum(p["monto"] for p in gestor.obtener_pagos()),
        notifications=gestor.obtener_notificaciones()[:5],
    )


@bp.route("/habitaciones")
def habitaciones():
    gestor = hotel_service.gestor
    piso = request.args.get("piso", "Todos")
    tipo = request.args.get("tipo", "Todos")
    estado = request.args.get("estado", "Todos")

    pisos = {}
    for numero_piso in range(1, 6):
        habitaciones_piso = [
            h for h in gestor.consultar_matriz().values()
            if str(h.piso) == str(numero_piso)
        ]
        pisos[numero_piso] = {
            "total": len(habitaciones_piso),
            "disponibles": sum(1 for h in habitaciones_piso if h.estado == "Disponible"),
            "ocupadas": sum(1 for h in habitaciones_piso if h.estado == "Ocupado"),
            "reservas": sum(1 for h in habitaciones_piso if h.estado == "Reserva"),
            "limpieza": sum(1 for h in habitaciones_piso if h.estado == "En limpieza"),
            "mantenimiento": sum(1 for h in habitaciones_piso if h.estado == "Mantenimiento"),
        }

    piso_seleccionado = piso if piso.isdigit() and 1 <= int(piso) <= 5 else None
    filtros_activos = tipo != "Todos" or estado != "Todos" or piso_seleccionado is not None
    habitaciones = (
        gestor.consultar_matriz_filtrada(piso, tipo, estado)
        if filtros_activos
        else []
    )
    return render_template(
        "habitaciones.html",
        habitaciones=habitaciones,
        pisos=pisos,
        piso_seleccionado=piso_seleccionado,
        filtros_activos=filtros_activos,
        piso=piso,
        tipo=tipo,
        estado=estado,
        motivos=gestor.motivos_mantenimiento,
    )


@bp.route("/habitaciones/<id_habitacion>", methods=["POST"])
def accion_habitacion(id_habitacion):
    accion = request.form.get("accion")
    gestor = hotel_service.gestor
    if accion == "mantenimiento":
        ejecutar_accion(gestor.poner_en_mantenimiento, id_habitacion, request.form.get("motivo", ""))
    elif accion == "finalizar_mantenimiento":
        ejecutar_accion(gestor.quitar_mantenimiento, id_habitacion, True)
    elif accion == "asignar_limpieza":
        ejecutar_accion(gestor.asignar_limpieza, id_habitacion, request.form.get("empleado", "Empleado 1"))
    elif accion == "finalizar_limpieza":
        ejecutar_accion(gestor.finalizar_limpieza, id_habitacion)
    elif accion == "reasignar":
        ejecutar_accion(gestor.reasignar_habitacion, id_habitacion, request.form.get("destino", ""))
    return redirect(url_for(
        "main.habitaciones",
        piso=request.args.get("piso", "Todos"),
        tipo=request.args.get("tipo", "Todos"),
        estado=request.args.get("estado", "Todos"),
    ))


@bp.route("/checkin", methods=["GET", "POST"])
def checkin():
    gestor = hotel_service.gestor
    ahora = hora_local()
    requiere_autorizacion = ahora.hour < 8 or ahora.hour >= 10

    if request.method == "POST":
        autorizacion_administrativa = request.form.get("administrador") == "1"

        if requiere_autorizacion and not autorizacion_administrativa:
            if ahora.hour < 8:
                flash("Antes de las 08:00 debe marcar la autorización administrativa e ingresar la contraseña del administrador.", "error")
            else:
                flash("Desde las 10:00 debe marcar la autorización administrativa e ingresar la contraseña del administrador.", "error")
            return redirect(url_for("main.checkin"))

        autorizado = not requiere_autorizacion
        if autorizacion_administrativa:
            password_ingresada = request.form.get("password_admin", "").strip()
            autorizado = validar_password_admin(password_ingresada)
            if not autorizado:
                flash("Contraseña administrativa incorrecta.", "error")
                return redirect(url_for("main.checkin"))

        ejecutar_accion(
            gestor.procesar_checkin,
            request.form.get("habitacion", ""),
            request.form.get("nombre", "").strip(),
            request.form.get("documento", "").strip(),
            ahora.hour,
            autorizado,
        )
        return redirect(url_for("main.checkin"))

    disponibles = [h for h in gestor.consultar_matriz().values() if h.estado in ("Disponible", "Reserva")]
    return render_template(
        "checkin.html",
        habitaciones=disponibles,
        habitacion_seleccionada=request.args.get("habitacion", ""),
        requiere_autorizacion=requiere_autorizacion,
        hora_actual=f"{ahora.hour % 12 or 12}:{ahora.minute:02d} {'a.m.' if ahora.hour < 12 else 'p.m.'}",
    )


@bp.route("/checkout", methods=["GET", "POST"])
def checkout():
    gestor = hotel_service.gestor
    if request.method == "POST":
        accion = request.form.get("accion", "checkout")
        id_habitacion = request.form.get("habitacion", "")
        if accion == "extender":
            try:
                noches = int(request.form.get("noches", "0"))
            except ValueError:
                noches = 0
            ejecutar_accion(gestor.extender_estancia, id_habitacion, noches)
        else:
            ejecutar_accion(gestor.procesar_checkout, id_habitacion)
        return redirect(url_for("main.checkout", habitacion=id_habitacion))
    ocupadas = [h for h in gestor.consultar_matriz().values() if h.estado == "Ocupado"]
    return render_template(
        "checkout.html",
        habitaciones=ocupadas,
        habitacion_seleccionada=request.args.get("habitacion", ""),
    )


@bp.route("/reservas", methods=["GET", "POST"])
def reservas():
    gestor = hotel_service.gestor
    if request.method == "POST":
        accion = request.form.get("accion")
        if accion == "crear":
            ejecutar_accion(
                gestor.crear_reserva,
                request.form.get("habitacion", ""),
                request.form.get("nombre", "").strip(),
                request.form.get("documento", "").strip(),
                request.form.get("fecha_entrada", "").strip(),
                request.form.get("fecha_salida", "").strip(),
            )
        elif accion == "cancelar":
            ejecutar_accion(gestor.cancelar_reserva, request.form.get("habitacion", ""))
        elif accion == "modificar":
            ejecutar_accion(
                gestor.modificar_fechas_reserva,
                request.form.get("habitacion", ""),
                request.form.get("fecha_entrada", ""),
                request.form.get("fecha_salida", ""),
            )
        elif accion == "noshow":
            ejecutar_accion(gestor.procesar_no_show)
        return redirect(url_for("main.reservas", habitacion=request.form.get("habitacion", "")))

    disponibles = [h for h in gestor.consultar_matriz().values() if h.estado in ("Disponible", "Reserva")]
    return render_template(
        "reservas.html",
        habitaciones=disponibles,
        habitacion_seleccionada=request.args.get("habitacion", ""),
        reservas=gestor.obtener_historial_reservas(),
    )


@bp.route("/limpieza", methods=["GET", "POST"])
def limpieza():
    gestor = hotel_service.gestor
    if request.method == "POST":
        accion = request.form.get("accion")
        id_habitacion = request.form.get("habitacion", "")
        if accion == "asignar":
            ejecutar_accion(gestor.asignar_limpieza, id_habitacion, request.form.get("empleado", "Empleado 1"))
        elif accion == "automatico":
            ejecutar_accion(gestor.asignar_limpieza_automatica)
        elif accion == "finalizar":
            ejecutar_accion(gestor.finalizar_limpieza, id_habitacion)
        return redirect(url_for("main.limpieza"))

    en_limpieza = [h for h in gestor.consultar_matriz().values() if h.estado in ("En limpieza", "Pendiente")]
    return render_template(
        "limpieza.html",
        habitaciones=en_limpieza,
        habitacion_seleccionada=request.args.get("habitacion", ""),
        empleados=gestor.empleados_limpieza,
        hoja=gestor.hoja_trabajo_limpieza(),
    )


@bp.route("/mantenimiento", methods=["GET", "POST"])
def mantenimiento():
    gestor = hotel_service.gestor
    if request.method == "POST":
        accion = request.form.get("accion")
        id_habitacion = request.form.get("habitacion", "")
        if accion == "poner":
            ejecutar_accion(gestor.poner_en_mantenimiento, id_habitacion, request.form.get("motivo", ""))
        elif accion == "finalizar":
            ejecutar_accion(gestor.quitar_mantenimiento, id_habitacion, True)
        return redirect(url_for("main.mantenimiento", habitacion=id_habitacion))

    habitaciones = list(gestor.consultar_matriz().values())
    return render_template(
        "mantenimiento.html",
        habitaciones=habitaciones,
        motivos=gestor.motivos_mantenimiento,
        habitacion_seleccionada=request.args.get("habitacion", ""),
    )


@bp.route("/reservas/calendario")
def calendario_reservas():
    gestor = hotel_service.gestor
    dias = request.args.get("dias", "30")
    try:
        dias = int(dias)
    except ValueError:
        dias = 30
    dias = dias if dias in (30, 60, 90) else 30
    return render_template(
        "calendario_reservas.html",
        dias=dias,
        proyeccion=gestor.proyectar_disponibilidad(dias),
        historial=gestor.obtener_historial_reservas(),
    )


@bp.route("/reservas/<id_habitacion>/comprobante")
def comprobante_reserva(id_habitacion):
    gestor = hotel_service.gestor
    ok, contenido = gestor.generar_comprobante_reserva(id_habitacion)
    if not ok:
        flash(contenido, "error")
        return redirect(url_for("main.reservas"))
    return Response(
        contenido,
        mimetype="text/plain",
        headers={"Content-Disposition": f"attachment; filename=reserva-{id_habitacion}.txt"},
    )


@bp.route("/buscar", methods=["GET"])
def buscar_huesped():
    criterio = request.args.get("q", "").strip()
    gestor = hotel_service.gestor
    resultados = gestor.buscar_huesped(criterio) if criterio else []
    return render_template("buscar.html", criterio=criterio, resultados=resultados)


@bp.route("/administracion", methods=["GET", "POST"])
def administracion():
    gestor = hotel_service.gestor
    if request.method == "POST":
        accion = request.form.get("accion")
        if accion == "tarifa":
            ejecutar_accion(gestor.actualizar_tarifa, request.form.get("tipo", ""), request.form.get("precio", "0"))
        elif accion == "descuento":
            ejecutar_accion(gestor.configurar_descuento_frecuente, request.form.get("porcentaje", "10"))
        return redirect(url_for("main.administracion"))

    tarifas = {}
    for habitacion in gestor.consultar_matriz().values():
        tarifas.setdefault(habitacion.tipo, habitacion.precio)
    return render_template(
        "administracion.html",
        tarifas=tarifas,
        descuento=gestor.descuento_cliente_frecuente,
    )


@bp.route("/pagos", methods=["GET", "POST"])
def pagos():
    gestor = hotel_service.gestor
    if request.method == "POST":
        try:
            monto = float(request.form.get("monto", "0") or 0)
        except ValueError:
            monto = 0
        ejecutar_accion(
            gestor.registrar_pago,
            request.form.get("habitacion", ""),
            monto,
            request.form.get("metodo", "Efectivo"),
        )
        return redirect(url_for("main.pagos"))
    ocupadas = [h for h in gestor.consultar_matriz().values() if h.estado == "Ocupado"]
    habitacion_param = request.args.get("habitacion", "")
    noches_param = request.args.get("noches", "1")
    try:
        noches_param = int(noches_param)
    except ValueError:
        noches_param = 1
    calculo = None
    if habitacion_param and noches_param > 0:
        calculo = gestor.calcular_costo_estadia(habitacion_param, noches_param)
    return render_template(
        "pagos.html",
        habitaciones=ocupadas,
        pagos=gestor.obtener_pagos(),
        habitacion_seleccionada=habitacion_param,
        noches=noches_param,
        calculo=calculo,
    )


@bp.route("/pagos/<id_habitacion>/comprobante")
def comprobante_pago(id_habitacion):
    gestor = hotel_service.gestor
    ok, contenido = gestor.generar_comprobante_pago(id_habitacion)
    if not ok:
        flash(contenido, "error")
        return redirect(url_for("main.pagos"))
    return Response(
        contenido,
        mimetype="text/plain",
        headers={"Content-Disposition": f"attachment; filename=pago-{id_habitacion}.txt"},
    )


@bp.route("/clientes", methods=["GET", "POST"])
def clientes():
    gestor = hotel_service.gestor
    if request.method == "POST":
        accion = request.form.get("accion")
        documento = request.form.get("documento", "").strip()
        if accion == "marcar_negra":
            ejecutar_accion(gestor.marcar_lista_negra, documento, request.form.get("motivo", ""))
        elif accion == "quitar_negra":
            ejecutar_accion(gestor.quitar_lista_negra, documento)
        elif accion == "marcar_frecuente":
            try:
                visitas = int(request.form.get("visitas", "6") or 6)
            except ValueError:
                visitas = 6
            ejecutar_accion(gestor.marcar_cliente_frecuente, documento, visitas)
        return redirect(url_for("main.clientes"))

    return render_template(
        "clientes.html",
        lista_negra=gestor.obtener_lista_negra(),
        frecuentes=gestor.obtener_clientes_frecuentes(),
        descuento=gestor.descuento_cliente_frecuente,
    )


@bp.route("/reportes")
def reportes():
    gestor = hotel_service.gestor
    fecha_inicio = request.args.get("fecha_inicio")
    fecha_fin = request.args.get("fecha_fin")
    ingresos = None
    if fecha_inicio and fecha_fin:
        try:
            inicio = datetime.date.fromisoformat(fecha_inicio)
            fin = datetime.date.fromisoformat(fecha_fin)
            ingresos = gestor.reporte_ingresos(inicio, fin)
            ingresos["total_ingresos"] = ingresos["total"]
            ingresos["cantidad_pagos"] = len(ingresos["pagos"])
        except ValueError:
            flash("Las fechas del reporte no son válidas.", "error")

    ocupacion = gestor.reporte_ocupacion_diaria()
    ocupacion["disponibles"] = sum(
        1 for h in gestor.consultar_matriz().values() if h.estado == "Disponible"
    )
    return render_template(
        "reportes.html",
        ocupacion=ocupacion,
        ingresos=ingresos,
        resumen=gestor.obtener_resumen_indicadores(),
        auditoria=list(reversed(gestor.obtener_auditoria()))[:100],
        clientes=gestor.clientes_frecuentes,
        reservas=gestor.obtener_historial_reservas(),
        turnos=gestor.obtener_historial_turnos(),
        fecha_inicio=fecha_inicio,
        fecha_fin=fecha_fin,
    )


@bp.route("/turno", methods=["GET", "POST"])
def turno():
    gestor = hotel_service.gestor
    if request.method == "POST":
        accion = request.form.get("accion")
        if accion == "abrir":
            recepcionista = request.form.get("recepcionista", "").strip()
            if not recepcionista:
                recepcionista = session.get("recepcionista", "")
            ejecutar_accion(gestor.abrir_turno, recepcionista)
        elif accion == "cerrar":
            ejecutar_accion(gestor.cerrar_turno)
        return redirect(url_for("main.turno"))

    return render_template(
        "turno.html",
        turno=gestor.turno_actual,
        historial=gestor.obtener_historial_turnos(),
        arqueo=gestor.generar_arqueo(),
    )


@bp.route("/reportes/exportar/<formato>")
def exportar_reporte(formato):
    gestor = hotel_service.gestor
    auditoria = gestor.obtener_auditoria()
    clientes = gestor.clientes_frecuentes

    if formato == "excel":
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill
        except ImportError:
            flash("Excel requiere instalar openpyxl.", "error")
            return redirect(url_for("main.reportes"))

        libro = Workbook()
        hoja = libro.active
        hoja.title = "Registro general"
        hoja.append(["Fecha", "Usuario / recepcionista", "Acción", "Habitación", "Detalle"])
        for fila in auditoria:
            hoja.append([fila["fecha"], fila["usuario"], fila["accion"], fila["id_habitacion"] or "—", fila["detalle"]])
        for celda in hoja[1]:
            celda.font = Font(bold=True, color="FFFFFF")
            celda.fill = PatternFill("solid", fgColor="0F766E")
        hoja.freeze_panes = "A2"

        hoja_clientes = libro.create_sheet("Clientes y visitas")
        hoja_clientes.append(["Documento", "Visitas", "Tipo", "Descuento próxima visita"])
        for documento, visitas in sorted(clientes.items(), key=lambda item: item[0]):
            frecuente = visitas > 5
            hoja_clientes.append([documento, visitas, "Frecuente" if frecuente else "Regular", "10%" if frecuente else "0%"])
        for celda in hoja_clientes[1]:
            celda.font = Font(bold=True, color="FFFFFF")
            celda.fill = PatternFill("solid", fgColor="0F766E")

        hoja_pagos = libro.create_sheet("Pagos")
        hoja_pagos.append(["Habitación", "Huésped", "Documento", "Base", "Descuento", "Total", "Método", "Fecha"])
        for pago in gestor.obtener_pagos():
            hoja_pagos.append([pago["id_habitacion"], pago["huesped"], pago["doc"], pago["monto_base"], pago["descuento"], pago["monto"], pago["metodo"], pago["fecha"]])
        for hoja_temp in (hoja, hoja_clientes, hoja_pagos):
            for columna in hoja_temp.columns:
                hoja_temp.column_dimensions[columna[0].column_letter].width = 24

        salida = io.BytesIO()
        libro.save(salida)
        salida.seek(0)
        return Response(
            salida.getvalue(),
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": "attachment; filename=registro_hotel.xlsx"},
        )

    if formato == "pdf":
        try:
            from reportlab.lib import colors
            from reportlab.lib.pagesizes import landscape, letter
            from reportlab.lib.styles import getSampleStyleSheet
            from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
        except ImportError:
            flash("PDF requiere instalar reportlab.", "error")
            return redirect(url_for("main.reportes"))

        salida = io.BytesIO()
        documento = SimpleDocTemplate(salida, pagesize=landscape(letter))
        estilos = getSampleStyleSheet()
        contenido = [Paragraph("Registro general del hotel", estilos["Title"]), Spacer(1, 12)]
        filas = [["Fecha", "Recepcionista", "Acción", "Habitación", "Detalle"]]
        for fila in auditoria[-100:]:
            filas.append([fila["fecha"], fila["usuario"], fila["accion"], fila["id_habitacion"] or "—", fila["detalle"][:120]])
        tabla = Table(filas, repeatRows=1)
        tabla.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0f766e")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
            ("FONT", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("FONTSIZE", (0, 0), (-1, -1), 7),
        ]))
        contenido.append(tabla)
        contenido.append(Spacer(1, 16))
        contenido.append(Paragraph("Clientes y visitas", estilos["Heading2"]))
        filas_clientes = [["Documento", "Visitas", "Categoría", "Descuento"]]
        for documento_cliente, visitas in sorted(clientes.items(), key=lambda item: item[0]):
            frecuente = visitas > 5
            filas_clientes.append([documento_cliente, str(visitas), "Frecuente" if frecuente else "Regular", "10%" if frecuente else "0%"])
        tabla_clientes = Table(filas_clientes, repeatRows=1)
        tabla_clientes.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#334155")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
            ("FONT", (0, 0), (-1, 0), "Helvetica-Bold"),
        ]))
        contenido.append(tabla_clientes)
        documento.build(contenido)
        salida.seek(0)
        return Response(
            salida.getvalue(),
            mimetype="application/pdf",
            headers={"Content-Disposition": "attachment; filename=registro_hotel.pdf"},
        )

    flash("Formato de exportación no válido.", "error")
    return redirect(url_for("main.reportes"))
